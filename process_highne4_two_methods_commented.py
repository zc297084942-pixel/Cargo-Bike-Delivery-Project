
# 1. IMPORT REQUIRED Python LIBRARIES
# These libraries read XML/Excel/CSV files, perform numerical
# calculations and nearest-node matching, and create the output
# Excel workbook. No modeling logic is applied in this section.

#import math, re, zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from collections import defaultdict

import numpy as np
import pandas as pd
from scipy.spatial import cKDTree
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# 2. USER SETTINGS AND INPUT/OUTPUT FILE PATHS
# CAPACITY is the maximum number of deliveries that one cargo
# bike can carry in a single hub-to-hub cycle. The remaining
# variables identify the XML plan, directed OD matrix, network
# node file, and Excel output location.

CAPACITY = 14
XML_PATH = Path('HighNE4.xml')
OD_CSV = Path('od_network_distance_matrix.csv')
NETWORK_XLSX = Path('b_line_network_combined_linked_nodes.xlsx')
OUT_XLSX = Path('HighNE4_basic_scenario_and_cycle_level_insertion_results.xlsx')
#OUT_CODE = Path('process_highne4_two_methods.py')


# 3. LOAD NETWORK NODES AND BUILD A SPATIAL SEARCH INDEX
# Reads node_id, x, and y from the Nodes worksheet. A cKDTree is
# built so every XML activity coordinate can be matched quickly
# to its nearest actual network node.

def load_nodes(network_xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(network_xlsx, read_only=True, data_only=True)
    ws = wb['Nodes']
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h:i for i, h in enumerate(headers)}
    node_ids=[]; coords=[]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx['node_id']] is None:
            continue
        node_ids.append(int(row[idx['node_id']]))
        coords.append((float(row[idx['x']]), float(row[idx['y']])))
    coords = np.array(coords, dtype=float)
    return np.array(node_ids, dtype=int), coords, cKDTree(coords)



# 4. LOAD THE DIRECTED ORIGIN-DESTINATION DISTANCE MATRIX
# The CSV is converted to a NumPy array. Later lookups follow:
# row = origin node, column = destination node.

def load_od(od_csv):
    df = pd.read_csv(od_csv, index_col=0)
    return df.to_numpy(dtype=float)



# 5. SELECT THE PLAN TO ANALYZE FOR ONE PERSON/SCENARIO
# MATSim may store multiple plans for a person. This function
# uses the plan marked selected="yes"; if none is marked, it
# falls back to the first available plan.

def chosen_plan(person):
    plans = person.findall('plan')
    for p in plans:
        if str(p.get('selected')).lower() == 'yes':
            return p
    return plans[0] if plans else None



# 6. PARSE THE XML INTO SCENARIOS, DELIVERIES, AND ORIGINAL CYCLES
# Each person is treated as one scenario. Delivery activities are
# numbered in XML order. Repeated micro_hub activities define the
# original cycle boundaries used by the cycle-level method and by
# the before-distance calculation.

def parse_xml_with_cycles(xml_path):
    root = ET.parse(xml_path).getroot()
    scenarios=[]
    for person in root.findall('person'):
        pid = person.get('id')
        plan = chosen_plan(person)
        if plan is None:
            continue
        act_counter=0
        hub=None
        deliveries=[]
        cycles=[]
        current=[]
        selected = plan.get('selected')
        for elem in list(plan):
            if elem.tag not in ('act', 'activity'):
                continue
            act_counter += 1
            typ = elem.get('type')
            x = float(elem.get('x')) if elem.get('x') is not None else None
            y = float(elem.get('y')) if elem.get('y') is not None else None
            if typ == 'micro_hub':
                if hub is None:
                    hub = {'x':x, 'y':y, 'display_label':1, 'act_no':act_counter}
                if current:
                    cycles.append(current)
                    current=[]
            elif typ == 'delivery':
                d = {
                    'person_id':pid,
                    'delivery_index':len(deliveries)+1,
                    'display_label':act_counter,
                    'x':x, 'y':y,
                    'start_time': elem.get('start_time',''),
                    'max_dur': elem.get('max_dur','') or elem.get('dur','') or '',
                }
                deliveries.append(d)
                current.append(d)
        if current:
            cycles.append(current)
        if hub is not None and deliveries:
            scenarios.append({'person_id':pid, 'selected_plan':selected, 'hub':hub, 'deliveries':deliveries, 'xml_cycles':cycles})
    return scenarios



# 7. MATCH ONE ACTIVITY COORDINATE TO THE NEAREST NETWORK NODE
# Returns the actual node ID and the coordinate-matching error.

def match_node(x,y,node_ids,coords,tree):
    dist, pos = tree.query([x,y], k=1)
    return int(node_ids[pos]), float(dist)



# 8. READ ONE DIRECTED OD DISTANCE AND CONVERT METERS TO KM
# Node IDs are treated as 1-based indices, so Python array indices
# use node_id - 1. Direction is preserved: a -> b.

def dist_km(a,b,od):
    return float(od[int(a)-1, int(b)-1]) / 1000.0


# 9. CALCULATE THE FULL CLOSED-ROUTE DISTANCE
# Adds hub -> first delivery, every delivery-to-delivery leg, and
# last delivery -> hub. This ensures every bike cycle returns to
# its micro-hub in the reported total distance.

def route_cost_km(hub_node, deliv_nodes, od):
    if not deliv_nodes:
        return 0.0
    total = dist_km(hub_node, deliv_nodes[0], od)
    for u,v in zip(deliv_nodes[:-1], deliv_nodes[1:]):
        total += dist_km(u,v,od)
    total += dist_km(deliv_nodes[-1], hub_node, od)
    return total



# 10. BASIC CHEAPEST-INSERTION HEURISTIC
# Starts with an empty hub-to-hub route. At each iteration, every
# unassigned delivery is tested at every possible route position.
# The delivery-position pair with the smallest added directed
# distance is inserted. This is the common heuristic used by both
# scenario-level and cycle-level methods.

def cheapest_insertion(hub_node, deliveries, od):
    remaining = deliveries[:]
    route = []
    while remaining:
        best = None
        for cand in remaining:
            cn = cand['node_id']
            for pos in range(len(route)+1):
                prev = hub_node if pos == 0 else route[pos-1]['node_id']
                nxt = hub_node if pos == len(route) else route[pos]['node_id']
                inc = dist_km(prev, cn, od) + dist_km(cn, nxt, od) - dist_km(prev, nxt, od)
                key = (inc, cand['delivery_index'], pos)
                if best is None or key < best[0]:
                    best = (key, cand, pos)
        _, chosen, pos = best
        route.insert(pos, chosen)
        remaining.remove(chosen)
    return route



# 11. SPLIT A SCENARIO-LEVEL SEQUENCE BY FIXED BIKE CAPACITY
# The first 14 deliveries become Cycle 1, the next 14 become Cycle
# 2, and so on. This is the non-hub-aware splitting step because
# the cut locations are based only on count, not hub-return cost.

def split_fixed_capacity(order, capacity=CAPACITY):
    return [order[i:i+capacity] for i in range(0, len(order), capacity)]



# 12. FORMAT ROUTES FOR REPORTING
# display_route uses the user-facing labels, where every hub is 1.
# actual_route preserves the matched network node IDs used in the
# OD matrix calculations.

def display_route(hub_node, cycle):
    return ' → '.join(map(str, [1] + [d['display_label'] for d in cycle] + [1]))


def actual_route(hub_node, cycle):
    return ' → '.join(map(str, [hub_node] + [d['node_id'] for d in cycle] + [hub_node]))



# 13. CREATE LEG-BY-LEG OUTPUT RECORDS
# Breaks each closed route into directed legs and records both the
# display labels and actual network nodes with each leg distance.

def leg_rows(file_name, scenario_id, method, cycle_no, route_type, hub_node, cycle, od):
    disp = [1] + [d['display_label'] for d in cycle] + [1]
    nodes = [hub_node] + [d['node_id'] for d in cycle] + [hub_node]
    rows=[]
    for i, (do,dd,no,nd) in enumerate(zip(disp[:-1], disp[1:], nodes[:-1], nodes[1:]), start=1):
        rows.append({
            'file':file_name, 'scenario_id':scenario_id, 'method':method, 'cycle_no':cycle_no, 'route_type':route_type,
            'leg_seq':i, 'from_display_node':do, 'to_display_node':dd,
            'from_actual_node':no, 'to_actual_node':nd,
            'distance_km':dist_km(no, nd, od)
        })
    return rows



# 14. RUN THE COMPLETE TWO-METHOD SIMULATION
# This function loads data, matches nodes, calculates the original
# XML distance, runs both insertion methods, and prepares all rows
# needed for the Excel output.

def process():
    # Load network nodes, the coordinate search tree, the directed OD
    # matrix, and all selected XML scenarios.
    node_ids, coords, tree = load_nodes(NETWORK_XLSX)
    od = load_od(OD_CSV)
    scenarios = parse_xml_with_cycles(XML_PATH)
    file_name = XML_PATH.name

    scenario_rows=[]; cycle_rows=[]; route_rows=[]; leg_details=[]; node_rows=[]; method_summary=defaultdict(lambda: {'scenarios':0,'deliveries':0,'cycles':0,'before_km':0.0,'after_km':0.0})
    max_match=0.0

    # Process one person/scenario at a time.
    for sc in scenarios:
        pid = sc['person_id']
        # Match the scenario's micro-hub coordinate to an actual node.
        hub_node, hub_err = match_node(sc['hub']['x'], sc['hub']['y'], node_ids, coords, tree)
        max_match = max(max_match, hub_err)
        # Match every delivery coordinate once and reuse the matched
        # node IDs in both insertion methods so the comparison is fair.
        # match deliveries once
        deliveries=[]
        for d in sc['deliveries']:
            node, err = match_node(d['x'], d['y'], node_ids, coords, tree)
            max_match = max(max_match, err)
            nd = dict(d); nd['node_id']=node; nd['match_error']=err
            deliveries.append(nd)
            node_rows.append({'file':file_name, 'scenario_id':pid, 'delivery_index':nd['delivery_index'], 'display_label':nd['display_label'], 'type':'delivery', 'x':nd['x'], 'y':nd['y'], 'actual_node':node, 'match_error':err})
        node_rows.append({'file':file_name, 'scenario_id':pid, 'delivery_index':'', 'display_label':1, 'type':'micro_hub', 'x':sc['hub']['x'], 'y':sc['hub']['y'], 'actual_node':hub_node, 'match_error':hub_err})
        # Reconstruct the original XML cycles using the matched delivery
        # records. These fixed groups are used by the cycle-level method.
        # remap cycle deliveries from original dicts to matched copies by delivery_index
        by_idx = {d['delivery_index']:d for d in deliveries}
        before_xml_cycles = [[by_idx[d['delivery_index']] for d in cyc] for cyc in sc['xml_cycles']]
        # Calculate the common baseline: sum the distance of every original
        # XML hub-to-hub cycle before applying either insertion method.
        before_km = sum(route_cost_km(hub_node, [d['node_id'] for d in cyc], od) for cyc in before_xml_cycles)


        # METHOD 1: BASIC SCENARIO-LEVEL INSERTION, NO HUB-AWARE SPLIT
        # Pool all deliveries from this scenario, run cheapest insertion
        # once across the full set, then cut the resulting sequence into
        # consecutive capacity-14 cycles. Deliveries may move across their
        # original XML cycle boundaries.

        # basic scenario-level insertion, no hub-aware split
        scen_order = cheapest_insertion(hub_node, deliveries, od)
        scen_after_cycles = split_fixed_capacity(scen_order, CAPACITY)
        methods = [('basic_scenario_level_no_hub_aware', scen_after_cycles)]
    
        # METHOD 2: CYCLE-LEVEL BASIC INSERTION
        # Keep each original XML cycle as a fixed delivery group and run
        # cheapest insertion independently within each group. Deliveries
        # cannot move from one original cycle to another.
       
        # cycle-level insertion within original cycles
        cycle_after_cycles = [cheapest_insertion(hub_node, cyc, od) for cyc in before_xml_cycles]
        methods.append(('cycle_level_basic_insertion', cycle_after_cycles))

        # Evaluate each method using the same directed OD matrix and the
        # same original before-distance baseline.
        for method, after_cycles in methods:
            after_km = sum(route_cost_km(hub_node, [d['node_id'] for d in cyc], od) for cyc in after_cycles)
            saved = before_km - after_km
            pct = saved / before_km if before_km else 0.0
            scenario_rows.append({'file':file_name,'scenario_id':pid,'selected_plan':sc['selected_plan'],'method':method,'deliveries':len(deliveries),'before_cycles':len(before_xml_cycles),'after_cycles':len(after_cycles),'capacity':CAPACITY,'hub_actual_node':hub_node,'before_km':before_km,'after_km':after_km,'saved_km':saved,'reduction_pct':pct})
            method_summary[method]['scenarios'] += 1
            method_summary[method]['deliveries'] += len(deliveries)
            method_summary[method]['cycles'] += len(after_cycles)
            method_summary[method]['before_km'] += before_km
            method_summary[method]['after_km'] += after_km

            # Save both original and optimized cycle routes so the Excel
            # workbook can show route-level and leg-level comparisons.
            # before cycles rows per method for direct comparison
            for route_type, cycles in [('before_original_xml_cycles', before_xml_cycles), ('after_'+method, after_cycles)]:
                for cno, cyc in enumerate(cycles, start=1):
                    cost = route_cost_km(hub_node, [d['node_id'] for d in cyc], od)
                    cycle_rows.append({'file':file_name,'scenario_id':pid,'method':method,'route_type':route_type,'cycle_no':cno,'stops':len(cyc),'distance_km':cost,'display_route':display_route(hub_node,cyc),'actual_node_route':actual_route(hub_node,cyc)})
                    route_rows.append({'file':file_name,'scenario_id':pid,'method':method,'route_type':route_type,'cycle_no':cno,'route_order_display':display_route(hub_node,cyc),'route_order_actual_nodes':actual_route(hub_node,cyc)})
                    leg_details.extend(leg_rows(file_name,pid,method,cno,route_type,hub_node,cyc,od))

    # Aggregate totals across all scenarios separately for each method.
    summary_rows=[]
    for method, v in method_summary.items():
        saved = v['before_km'] - v['after_km']
        pct = saved / v['before_km'] if v['before_km'] else 0.0
        summary_rows.append({'file':file_name,'method':method,'scenarios':v['scenarios'],'deliveries':v['deliveries'],'cycles':v['cycles'],'before_km':v['before_km'],'after_km':v['after_km'],'saved_km':saved,'reduction_pct':pct,'max_match_error':max_match})

    return summary_rows, scenario_rows, cycle_rows, route_rows, leg_details, node_rows



# 15. EXPORT ALL RESULTS TO A FORMATTED EXCEL WORKBOOK
# Creates method summaries, scenario and cycle comparisons, route
# orders, leg details, node matching, and explanatory notes.

def write_xlsx(summary_rows, scenario_rows, cycle_rows, route_rows, leg_rows, node_rows):
    wb = Workbook()
    # Rename the default sheet, then create one worksheet for each
    # level of output detail.
    # remove default
    ws = wb.active
    ws.title = 'Method_Summary'
    sheets = [
        ('Method_Summary', summary_rows),
        ('Scenario_Comparison', scenario_rows),
        ('Cycle_Comparison', cycle_rows),
        ('Route_Order_Display', route_rows),
        ('Leg_Details_Display', leg_rows),
        ('Node_Mapping', node_rows),
        ('Notes', [
            {'item':'Uploaded XML file','description':str(XML_PATH.name)},
            {'item':'Selected plan handling','description':'For each person/scenario, the selected="yes" plan is used; if no selected plan exists, the first plan is used.'},
            {'item':'OD matrix direction','description':'Directed distance lookup uses row = origin node and column = destination node.'},
            {'item':'Capacity','description':str(CAPACITY)},
            {'item':'basic_scenario_level_no_hub_aware','description':'Collect all deliveries in the scenario, run basic cheapest insertion once, then split the resulting sequence every 14 deliveries. Split points are not optimized for hub return distance.'},
            {'item':'cycle_level_basic_insertion','description':'Use original XML micro-hub rows as fixed cycle boundaries and run basic cheapest insertion separately within each cycle. Deliveries cannot move across original cycles.'},
            {'item':'Negative savings','description':'Possible because basic insertion is a greedy heuristic and no acceptance safeguard is applied.'},
        ])
    ]
    # Write dictionaries as worksheet rows, then apply consistent
    # headers, number formats, filters, frozen panes, and column widths.
    # Create/update sheets
    for idx, (name, rows) in enumerate(sheets):
        if idx == 0:
            ws = wb[name]
        else:
            ws = wb.create_sheet(name)
        if not rows:
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for r in rows:
            ws.append([r.get(h,'') for h in headers])
        # style
        header_fill = PatternFill('solid', fgColor='1F4E78')
        header_font = Font(bold=True, color='FFFFFF')
        thin = Side(style='thin', color='D9E2F3')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(bottom=thin)
                if isinstance(cell.value, float):
                    if 'pct' in str(ws.cell(row=1, column=cell.column).value):
                        cell.number_format = '0.00%'
                    else:
                        cell.number_format = '0.000'
                cell.alignment = Alignment(vertical='top', wrap_text=False)
        # widths
        for col_idx, col in enumerate(ws.columns, start=1):
            header = str(ws.cell(row=1, column=col_idx).value)
            max_len = len(header)
            sample_count = 0
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 80))
                sample_count += 1
                if sample_count > 200:
                    break
            width = min(max(max_len + 2, 10), 55)
            if 'route' in header.lower(): width = 55
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    wb.save(OUT_XLSX)


# 16. PROGRAM ENTRY POINT
# Runs the simulation, writes the workbook, and prints the output
# path plus the method-level summary to the terminal.

if __name__ == '__main__':
    data = process()
    write_xlsx(*data)
    print(OUT_XLSX)
    for row in data[0]:
        print(row)
